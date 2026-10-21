"""What a workbook cell holds, against any text the pipeline can store.

Every string on the Questions and Citations sheets is a document's, a
model's or a person's: a topic name typed on the Topics page or answered in
Argilla, a statement or an answer a model wrote. PostgreSQL text holds any
character but NUL, so the workbook meets every one of them.
"""

from __future__ import annotations

import io

import pytest
from hypothesis import given
from hypothesis import strategies as st
from openpyxl import load_workbook
from test_export import cited, stored

from question_generation.export import _CELL_LIMIT, workbook

#: Text an XML 1.0 document can carry, which is what an .xlsx is: no C0
#: control but tab, newline and return, no surrogate, no U+FFFE or U+FFFF.
_LEGAL = st.text(
    alphabet=st.characters(
        blacklist_categories=("Cs",),
        blacklist_characters=[
            *(chr(code) for code in range(0x20) if code not in (0x09, 0x0A, 0x0D)),
            "￾",
            "￿",
        ],
    ),
    min_size=1,
    max_size=200,
)


def _cell(drawn: bytes, sheet: str, heading: str):
    """One cell under a heading on the first data row, read the way Excel does."""
    book = load_workbook(io.BytesIO(drawn))
    rows = list(book[sheet].iter_rows(values_only=True))
    return rows[1][list(rows[0]).index(heading)]


@given(text=_LEGAL)
def test_any_legal_text_reads_back_as_written(text) -> None:
    """German, English, a third script, emoji or a leading `=`: what was stored."""
    drawn = workbook(
        [stored(question_text=text, topics=[text])], [cited(evidence_text=text)]
    )

    assert _cell(drawn, "Questions", "Question") == text
    assert _cell(drawn, "Questions", "Topics") == text
    assert _cell(drawn, "Citations", "Evidence") == text


def test_a_cell_over_the_limit_is_cut_to_the_limit_and_marked() -> None:
    """Still under Excel's 32,767 once the mark is added."""
    drawn = workbook([stored(target_answer="ä" * (_CELL_LIMIT + 5))])
    cell = _cell(drawn, "Questions", "Answer")

    assert cell.startswith("ä" * _CELL_LIMIT)
    assert cell.endswith("[truncated]")
    assert len(cell) <= 32_767


@pytest.mark.parametrize(
    "where",
    [
        {"topics": ["Konto\x0bführung"]},
        {"target_answer": "Gebühr\x0c 2,50 EUR"},
        {"question_text": "When is the card\x1b blocked?"},
    ],
    ids=["topic-vertical-tab", "answer-form-feed", "question-escape"],
)
def test_a_control_character_does_not_stop_the_export(where) -> None:
    """The workbook is written, the character dropped, the rest of the text kept."""
    drawn = workbook([stored(**where)])

    (column, (value,)) = next(
        (key, value if isinstance(value, list) else [value])
        for key, value in where.items()
    )
    heading = {"topics": "Topics", "target_answer": "Answer"}.get(column, "Question")
    kept = "".join(ch for ch in value if ch >= " " or ch in "\t\n\r")
    assert _cell(drawn, "Questions", heading) == kept
