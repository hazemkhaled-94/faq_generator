"""The questions as a workbook.

A spreadsheet is the one artefact here that leaves the deployment and gets
read by somebody who has never seen the pipeline, so what is pinned is what
a reader would notice: that the sheets are there, that a list is a list of
words rather than Python's repr of one, and that nothing raises on the two
inputs that are not a full set - nothing selected, and a cell longer than
Excel will hold.
"""

from __future__ import annotations

import io
from contextlib import contextmanager

import pandas as pd
import pytest

from question_generation.export import (
    _CELL_LIMIT,
    NEVER_DOCUMENTS,
    NEVER_PASSAGES,
    _flatten,
    workbook,
)
from question_generation.models import (
    Citation,
    QuestionQuality,
    StoredQuestion,
    pages,
)
from question_generation.queue import NeverAsked


def stored(id: int = 1, **overrides) -> StoredQuestion:
    """One question, as the listing hands it over."""
    fields = {
        "id": id,
        "question_text": "Within how many hours is a request answered?",
        "target_answer": "48 hours",
        "answer_explanation": "A standard request is answered within 48 hours.",
        "answerable": True,
        "difficulty": "easy",
        "passage_scope": "single_passage",
        "document_scope": "single_document",
        "topic_scope": "single_topic",
        "answer_chars": 8,
        "language": "en",
        "status": "accepted",
        "rejected_reason": None,
        "created_at": "2026-09-24T12:00:00+00:00",
        "facts": 1,
        "documents": ["Service policy"],
        "topics": ["Support", "Policy"],
        "statements": ["A standard request is answered in 48 hours."],
        "evidence": ["Requests are answered in 48 hours."],
    }
    return StoredQuestion(**{**fields, **overrides})


def cited(question_id: int = 1, **overrides) -> Citation:
    """One fact a question rests on."""
    fields = {
        "question_id": question_id,
        "fact_id": 7,
        "statement": "A standard request is answered in 48 hours.",
        "evidence_text": "Requests are answered in 48 hours.",
        "validated": True,
        "document": "Service policy",
        "doc_sha256": "a1b2c3",
        "ordinal": 4,
        "page": 2,
    }
    return Citation(**{**fields, **overrides})


def sheets(drawn: bytes) -> dict[str, pd.DataFrame]:
    """Reads a built workbook back the way a recipient opens it."""
    book = pd.ExcelFile(io.BytesIO(drawn))
    return {name: book.parse(name) for name in book.sheet_names}


def test_the_workbook_has_the_four_sheets() -> None:
    """Questions, what they cite, what a judge said, and the set's shape."""
    read = sheets(workbook([stored()], [cited()]))

    assert list(read) == ["Questions", "Citations", "Assessment", "Summary"]
    assert len(read["Questions"]) == 1
    assert len(read["Citations"]) == 1


def test_a_question_carries_the_columns_a_reader_filters_on() -> None:
    """The headings are words, and the values are under them."""
    row = sheets(workbook([stored()]))["Questions"].iloc[0]

    assert row["Question"] == "Within how many hours is a request answered?"
    assert row["Answer"] == "48 hours"
    assert row["Type"] is not None
    assert row["Difficulty"] == "easy"
    assert row["Status"] == "accepted"


def test_a_list_is_joined_rather_than_repr_ed() -> None:
    """`topics` is a list, and a cell has no type for one.

    Left to `str()` it would reach a reader as `['Support', 'Policy']`,
    brackets and quotes included.
    """
    row = sheets(workbook([stored()]))["Questions"].iloc[0]

    assert row["Topics"] == "Support, Policy"


def test_a_question_row_carries_what_it_rests_on() -> None:
    """The facts, the evidence and the reason, on the question's own row.

    The Citations sheet holds the same pairs one row apiece, and a reader
    who has filtered the Questions sheet to the refused ones should not
    have to join it by id to see what any of them cited.
    """
    row = sheets(workbook([stored()]))["Questions"].iloc[0]

    assert row["The facts it cites"] == "A standard request is answered in 48 hours."
    assert row["The evidence they rest on"] == "Requests are answered in 48 hours."
    assert row["Answer in full"] == "A standard request is answered within 48 hours."
    # The title, not the digest: a sha256 is not a thing anybody can look
    # up by hand, which is why the Citations sheet already resolves one.
    assert row["Documents"] == "Service policy"


def test_a_question_says_whether_it_is_in_the_release() -> None:
    """The Released column, beside the status."""
    read = sheets(workbook([stored(1, released=True), stored(2)]))["Questions"]

    assert list(read["Released"]) == [True, False]
    assert list(read.columns).index("Released") == (
        list(read.columns).index("Status") + 1
    )


def test_a_question_names_the_categories_it_was_written_as() -> None:
    """As `categorization=category`, after its language; blank where it has none."""
    read = sheets(
        workbook([stored(1, categories=["Phrasing=concise-and-natural"]), stored(2)])
    )["Questions"]

    assert read.iloc[0]["Categories"] == "Phrasing=concise-and-natural"
    assert pd.isna(read.iloc[1]["Categories"])
    assert list(read.columns).index("Categories") == (
        list(read.columns).index("Language") + 1
    )


def test_sentences_are_one_per_line_rather_than_comma_joined() -> None:
    """A comma between two sentences reads as one sentence."""
    cited = ["The first holds.", "The second holds."]
    row = sheets(workbook([stored(statements=cited)]))["Questions"].iloc[0]

    assert row["The facts it cites"] == "The first holds.\nThe second holds."


class _Judged:
    """One question's judgement, as the assessment catalogue hands it over."""

    def __init__(self, *metrics) -> None:
        """Holds these judgements and nothing else."""
        self.metrics = metrics
        self.approved = all(one.approved for one in metrics)
        self.judge_model = "a-model"
        self.refused = [one.metric for one in metrics if not one.approved]
        self.disagrees = False


class _Metric:
    """One metric of one judgement."""

    def __init__(self, metric: str, label: str, approved: bool, why: str) -> None:
        """Holds what the judge said about this one metric."""
        self.metric, self.label, self.approved, self.explanation = (
            metric,
            label,
            approved,
            why,
        )
        self.score = 1.0 if approved else 0.0


def test_the_judges_reasoning_is_on_the_question_row() -> None:
    """Not only on the Assessment sheet.

    A reader sorting the questions by what the judge refused wants the
    reason in the row they are looking at; a label with no reasoning
    beside it is one nobody can check.
    """
    refused = _Metric("groundedness", "unsupported", False, "The passages say 24.")
    row = sheets(workbook([stored()], judged={1: _Judged(refused)}))["Questions"].iloc[
        0
    ]

    assert row["Judge"] == "refused"
    assert row["Judge refused by"] == "groundedness"
    assert "The passages say 24." in row["Why the judge said so"]


def test_an_unjudged_question_says_so_rather_than_going_blank() -> None:
    """A blank cell reads as "no problem found"."""
    row = sheets(workbook([stored()]))["Questions"].iloc[0]

    assert row["Judge"] == "not judged"
    assert pd.isna(row["Why the judge said so"])


def test_nothing_selected_still_has_its_headings() -> None:
    """An empty filter writes an empty sheet, not a broken one.

    A workbook of nothing is a workbook somebody can see the shape of; a
    raise here would mean the one filter that matches no question is the
    one that cannot be exported.
    """
    read = sheets(workbook([], []))

    assert list(read) == ["Questions", "Citations", "Assessment", "Summary"]
    assert read["Questions"].empty
    assert "Question" in read["Questions"].columns


@pytest.mark.parametrize(
    ("first", "last", "page"),
    [(3, 5, "3–5"), (2, 2, 2), (2, None, 2), (None, None, None)],
)
def test_a_citation_names_every_page_its_passage_runs_over(first, last, page) -> None:
    """The start alone where the passage stays on one page."""
    assert pages(first, last) == page


def test_a_cell_too_long_for_excel_is_cut_and_marked() -> None:
    """Excel refuses a cell over 32,767 characters outright.

    Cut with the cut named, so nobody reads a truncated answer as a short
    one - and so a workbook of thousands does not raise on its last row.
    """
    cut = _flatten("x" * (_CELL_LIMIT + 500))

    assert len(cut) < _CELL_LIMIT + 50
    assert cut.endswith("[truncated]")


def test_a_long_answer_survives_the_round_trip() -> None:
    """The cut is at Excel's limit, so an ordinary long answer is untouched."""
    answer = "a sentence. " * 200
    row = sheets(workbook([stored(target_answer=answer)]))["Questions"].iloc[0]

    assert row["Answer"] == answer


def test_text_that_starts_like_a_formula_is_written_as_text() -> None:
    """Text from a document or a model is never a live formula in the file.

    A string that starts with = is otherwise stored as a formula, which Excel
    runs when the recipient opens the workbook. Kept as typed, character for
    character, so the reader sees what the pipeline wrote.
    """
    from openpyxl import load_workbook

    answer = '=HYPERLINK("http://x","y")'
    book = load_workbook(
        io.BytesIO(
            workbook(
                [stored(target_answer=answer)],
                [cited(evidence_text="=1+1")],
            )
        )
    )

    found = [
        cell
        for sheet in book.worksheets
        for row in sheet.iter_rows(min_row=2)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert {cell.value for cell in found} == {answer, "=1+1"}
    assert {cell.data_type for cell in found} == {"s"}


@pytest.mark.parametrize("value", [None, 1, True, "plain"])
def test_an_ordinary_value_passes_through(value) -> None:
    """Only lists and over-long text are rewritten."""
    assert _flatten(value) == value


def test_the_summary_counts_what_the_workbook_holds() -> None:
    """Written so a workbook mailed on says what it is.

    Read off the quality report the page was showing rather than counted
    again here, so the figures in the file are the figures on screen.
    """
    quality = QuestionQuality(
        total=4,
        accepted=3,
        draft=0,
        unanswerable=1,
        followups=0,
        topics_covered=1,
        topics_in_coverage=2,
        mean_question_chars=40.0,
        mean_answer_chars=8.0,
        rejected={"duplicate": 1},
        difficulty={"easy": 3, "hard": 1},
        passage_scope={},
        document_scope={},
        topic_scope={},
        question_type={"factoid": 4},
        answer_form={},
        planned_difficulty={},
        cognitive_level={"recall": 4},
        cognitive_level_checked=4,
        planned_met=3,
    )
    summary = sheets(workbook([stored()], [cited()], quality))["Summary"]
    counted = dict(zip(summary["Measure"], summary["Count"], strict=True))

    assert counted["Questions in this workbook"] == 1
    assert counted["Citation rows"] == 1
    assert counted["Matching the filter"] == 4
    assert counted["difficulty: easy"] == 3
    assert counted["rejected: duplicate"] == 1


def test_the_summary_counts_the_kinds_and_the_categories() -> None:
    """The same counts the quality report holds, which the page shows too."""
    quality = QuestionQuality(
        *[0] * 9,
        *[{}] * 9,
        cognitive_level_checked=0,
        planned_met=0,
        kind={"single": 3, "compound": 1},
        categories={"Phrasing=concise-and-natural": 2},
    )
    summary = sheets(workbook([stored()], [cited()], quality))["Summary"]
    counted = dict(zip(summary["Measure"], summary["Count"], strict=True))

    assert counted["kind: compound"] == 1
    assert counted["categories: Phrasing=concise-and-natural"] == 2


def test_the_summary_names_the_models() -> None:
    """A tester judging their chatbot reads which models made the questions."""
    quality = QuestionQuality(
        *[0] * 9,
        *[{}] * 9,
        cognitive_level_checked=0,
        planned_met=0,
        writer_model={"azure/gpt-5.4": 2},
        verifier_model={"azure/gpt-4.1": 2},
        phrasing_model={"azure/gpt-4.1": 1, "not run or not recorded": 1},
    )
    summary = sheets(workbook([stored()], [cited()], quality))["Summary"]
    counted = dict(zip(summary["Measure"], summary["Count"], strict=True))

    assert counted["writer_model: azure/gpt-5.4"] == 2
    assert counted["verifier_model: azure/gpt-4.1"] == 2
    assert counted["phrasing_model: azure/gpt-4.1"] == 1
    assert counted["phrasing_model: not run or not recorded"] == 1


class _Catalog:
    """A catalogue that records how the export asked for its rows."""

    def __init__(self, rows: int = 3) -> None:
        """Holds that many questions."""
        self._rows = rows
        self.limits: list[int | None] = []
        self.asked: dict[str, dict] = {}

    @contextmanager
    def snapshot(self):
        """The one session the three reads share, which here is nothing."""
        yield None

    def page(self, limit=50, read=None, **where):
        """Records the window asked for and answers within it."""
        self.asked["page"] = where
        self.limits.append(limit)
        taken = self._rows if limit is None else min(limit, self._rows)
        return self._rows, [stored(one) for one in range(1, taken + 1)]

    def citations(self, read=None, **where):
        """No citations."""
        self.asked["citations"] = where
        return []

    def quality(self, read=None, **where):
        """No report, which the workbook takes as None."""
        self.asked["quality"] = where


class _Nothing:
    """A question queue that reaches no database for what no question reaches."""

    @staticmethod
    def never_asked() -> NeverAsked:
        """Every document has been asked."""
        return NeverAsked((), 0)


class _NoJudgements:
    """An assessment catalogue for a corpus nobody has judged.

    Which is every corpus until the evaluation phase is switched on, so
    this is the ordinary case rather than a stub for an unusual one.
    """

    def verdicts_for(self, kind: str, ids: list[int]) -> dict:
        """Nothing has been judged."""
        del kind, ids
        return {}


def test_the_export_asks_for_every_row_in_one_window(tmp_path, monkeypatch) -> None:
    """It counted, then asked for a window the size of the count.

    Two statements with a gap between them: a question written into the
    gap did not fit the window the count had already decided on, so the
    workbook came up a row short of the figure on its own Summary sheet.

    Patched where `main` imports from rather than on the module: those
    three imports are inside the function, so the name this binds is the
    one it looks up when it runs.
    """
    from question_generation import export as module

    catalog = _Catalog(rows=3)
    monkeypatch.setattr("question_generation.catalog.QuestionCatalog", lambda: catalog)
    monkeypatch.setattr("database.qa_generator.engine", lambda: None)
    monkeypatch.setattr("telemetry.trace_engine", lambda engine: None)
    monkeypatch.setattr("question_generation.export.QuestionQueue", _Nothing)
    # The judge's verdicts are looked up the same way the questions are,
    # and are the same kind of import inside `main`. Patched here so this
    # stays a test about the window and not about a database.
    monkeypatch.setattr(
        "assessment.repository.AssessmentCatalog", lambda: _NoJudgements()
    )
    written = tmp_path / "questions.xlsx"

    module.main(["--out", str(written), "--no-citations"])

    assert catalog.limits == [None], "one unlimited read, not a count and a window"
    assert written.exists()


@pytest.mark.parametrize(
    ("flags", "released"), [(["--released"], True), ([], None)], ids=["on", "off"]
)
def test_released_narrows_every_read_to_the_release(
    tmp_path, monkeypatch, flags, released
) -> None:
    """The rows, their citations and the Summary, combined with the rest."""
    from question_generation import export as module

    catalog = _Catalog(rows=1)
    monkeypatch.setattr("question_generation.catalog.QuestionCatalog", lambda: catalog)
    monkeypatch.setattr("database.qa_generator.engine", lambda: None)
    monkeypatch.setattr("telemetry.trace_engine", lambda engine: None)
    monkeypatch.setattr("question_generation.export.QuestionQueue", _Nothing)
    monkeypatch.setattr(
        "assessment.repository.AssessmentCatalog", lambda: _NoJudgements()
    )

    module.main(
        ["--out", str(tmp_path / "release.xlsx"), "--status", "accepted", *flags]
    )

    assert set(catalog.asked) == {"page", "citations", "quality"}
    for read, where in catalog.asked.items():
        assert where["released"] is released, read
        assert where["status"] == "accepted", read


def test_an_unanswerable_row_shows_its_behaviour() -> None:
    """What a good reply does for its kind stands in its Answer cell.

    An answerable row keeps its key; an unanswerable one stored with no kind
    is blank.
    """
    read = sheets(
        workbook(
            [
                stored(
                    1,
                    answerable=False,
                    target_answer=None,
                    categories=["Unanswerable=false-presupposition"],
                    behaviour="points out ...",
                ),
                stored(2),
                stored(3, answerable=False, target_answer=None),
            ]
        )
    )["Questions"]

    assert read.iloc[0]["Answer"] == "Expected behaviour: points out ..."
    assert read.iloc[1]["Answer"] == "48 hours"
    assert pd.isna(read.iloc[2]["Answer"])


def test_the_catalog_reads_the_behaviour_off_the_stored_categories() -> None:
    """From the stored category that carries it, and None where none does."""
    from collections import namedtuple
    from dataclasses import asdict

    from question_generation.catalog import _stored

    kind = {
        "categorization": "Unanswerable",
        "category": "false-presupposition",
        "role": "question",
        "form": "question",
        "description": "d.",
        "settles": "refutes",
        "behaviour": "points out ...",
    }
    columns = {
        key: value for key, value in asdict(stored()).items() if key != "behaviour"
    }
    row = namedtuple("Row", list(columns))

    def read(categories):
        return _stored(row(**columns | {"created_at": None, "categories": categories}))

    assert read([kind]).behaviour == "points out ..."
    assert read([kind]).categories == ["Unanswerable=false-presupposition"]
    assert read(None).behaviour is None


def test_two_documents_sharing_a_title_are_told_apart_by_their_id() -> None:
    """Both sheets carry the document's digest head, which a title does not."""
    first, second = "a1" * 32, "b2" * 32
    read = sheets(
        workbook(
            [
                stored(1, document_ids=[first[:12]]),
                stored(2, document_ids=[second[:12]]),
            ],
            [
                cited(1, document="Service policy", doc_sha256=first),
                cited(2, document="Service policy", doc_sha256=second),
            ],
        )
    )

    assert list(read["Questions"]["Document IDs"]) == [first[:12], second[:12]]
    assert list(read["Citations"]["Document ID"]) == [first[:12], second[:12]]


# ── What no accepted question reaches ─────────────────────────────────────


@pytest.mark.parametrize(
    "never",
    [
        NeverAsked((("Terms of sale", 12), ("Allgemeine Geschäftsbedingungen", 4)), 31),
        NeverAsked((), 0),
    ],
    ids=["documents named", "none"],
)
def test_the_summary_counts_the_documents_and_passages_no_question_reaches(
    never,
) -> None:
    """The corpus, beside what the filter selected: a reader sees what is missing."""
    summary = sheets(workbook([stored()], [cited()], None, None, never))["Summary"]
    counted = dict(zip(summary["Measure"], summary["Count"], strict=True))

    assert counted[NEVER_DOCUMENTS] == len(never.documents)
    assert counted[NEVER_PASSAGES] == never.passages
    for name, held in never.documents:
        assert counted[f"{NEVER_DOCUMENTS}: {name}"] == held


def test_a_workbook_with_no_coverage_says_nothing_of_it() -> None:
    """The api and the command line hand it over; a caller that has none leaves it out."""
    summary = sheets(workbook([stored()], [cited()]))["Summary"]

    assert NEVER_DOCUMENTS not in set(summary["Measure"])


def test_the_status_line_names_the_documents_with_their_passages() -> None:
    """Counts and names, in whichever language the titles are in."""
    never = NeverAsked((("Terms of sale", 12), ("Kreditbedingungen", 4)), 31)

    assert str(never) == (
        "2 document(s) with facts and no accepted question: Terms of sale (12), "
        "Kreditbedingungen (4); 31 passage(s) with facts never asked"
    )
    assert str(NeverAsked((), 3)) == (
        "every document with facts has an accepted question; "
        "3 passage(s) with facts never asked"
    )


def test_the_status_line_stops_naming_after_twenty() -> None:
    """A corpus ten times the size does not print its whole contents."""
    never = NeverAsked(tuple((f"doc-{n}", 1) for n in range(25)), 25)

    assert "doc-19 (1) and 5 more;" in str(never)
    assert "doc-20" not in str(never)


def test_diversity_by_document_prints_what_the_summary_says_no_question_reaches(
    tmp_path, capsys
) -> None:
    """The curve covers the documents that have questions; this is the other half."""
    from assessment import run as assessment_run

    never = NeverAsked((("Terms of sale", 12),), 14)
    path = tmp_path / "set.xlsx"
    path.write_bytes(workbook([stored()], [cited()], None, None, never))

    assessment_run._never_asked(path)

    said = capsys.readouterr().out
    assert f"{NEVER_DOCUMENTS}: 1" in said
    assert f"{NEVER_PASSAGES}: 14" in said
    assert "Terms of sale (12 passage(s))" in said


def test_the_two_modules_read_the_same_summary_words() -> None:
    """No service imports another, so the words are written twice and held here."""
    from assessment import run as assessment_run

    assert assessment_run._NEVER == (NEVER_DOCUMENTS, NEVER_PASSAGES)
