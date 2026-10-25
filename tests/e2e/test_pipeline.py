"""One document, all the way through, in this process.

Every stage is the real service against the real database and the real
object store. Four things are stood in for: the PDF converter, whose layout
models are a download; the served models, which are a GPU somewhere; the
embedding model, which is 2.2 GB of weights and which the dedup gate needs
only to be consistent; and the topic fit, which is the one stood in for
because it cannot run here rather than because it is expensive - the
frequency filter needs more documents than one fixture has.

What this covers that the per-stage tests do not is the handover: each
stage leaves a row in the state the next one claims, and nothing starts
until something asks it to.
"""

from __future__ import annotations

import re
from typing import ClassVar

import pytest
import weights
from sqlalchemy import text

from database.qa_generator import Status
from extraction.extractors import ExtractorRegistry
from extraction.extractors.llm import LlmExtractor
from extraction.extractors.table import TableExtractor
from extraction.repository import PassageQueue
from extraction.service import ExtractionService
from extraction.validation import FactChecker
from ingestion.models import UploadedFile
from ingestion.repository import DocumentRepository
from ingestion.service import IngestService
from preprocessing.chunking.repository import ChunkQueue
from preprocessing.chunking.service import ChunkingService
from preprocessing.parsing.analysis import DocumentAnalyser
from preprocessing.parsing.models import Conversion, SourceDocument
from preprocessing.parsing.pipelines import Pipeline, PipelineRegistry
from preprocessing.parsing.repository import ParseQueue
from preprocessing.parsing.service import ParsingService
from question_generation.catalog import QuestionCatalog
from question_generation.checker import QuestionChecker
from question_generation.config import Settings as QuestionSettings
from question_generation.generation import QuestionWriter
from question_generation.queue import QuestionQueue
from question_generation.service import QuestionGenerationService, reverify
from question_generation.verifier import Verifier

pytestmark = [pytest.mark.integration, pytest.mark.e2e]

#: What the fake converter produces: two English paragraphs carrying claims
#: a deterministic reader will not take, so the model is what reads them.
PARAGRAPHS = (
    "The device weighs 4 kg and runs for 12 hours.",
    "The device arrives in March 2026 from the Hamburg plant.",
)

PDF = (
    b"%PDF-1.4\n"
    b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
    b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\n"
    b"trailer<</Root 1 0 R>>\n"
)


#: Two short documents of a shop's terms, one German and one English, marked
#: by a comment in their file. Each section is one passage of one sentence,
#: and the statement the stub reads out of it is the sentence's first clause:
#: two sections under one title are what the bridge pass joins.
SHOPS = {
    b"%shop-en": {
        "title": "Terms of the shop",
        "sections": (
            (
                "Shipping",
                "A parcel leaves the warehouse within 2 working days of payment and is tracked online.",
                "A parcel leaves the warehouse within 2 working days of payment.",
            ),
            (
                "Returns",
                (
                    "A returned item is refunded within 14 days of its arrival and the "
                    "refund goes back to the card."
                ),
                "A returned item is refunded within 14 days of its arrival.",
            ),
        ),
    },
    b"%shop-de": {
        "title": "Bedingungen des Shops",
        "sections": (
            (
                "Versand",
                (
                    "Eine Bestellung wird innerhalb von 2 Werktagen nach der Zahlung "
                    "versandt und kann online verfolgt werden."
                ),
                (
                    "Eine Bestellung wird innerhalb von 2 Werktagen nach der Zahlung "
                    "versandt."
                ),
            ),
            (
                "Rücksendung",
                (
                    "Eine zurückgesendete Ware wird innerhalb von 14 Tagen nach ihrem "
                    "Eingang erstattet und der Betrag geht auf die Karte zurück."
                ),
                (
                    "Eine zurückgesendete Ware wird innerhalb von 14 Tagen nach ihrem "
                    "Eingang erstattet."
                ),
            ),
        ),
    },
}

#: What the stub reads out of each shop sentence.
CLAUSES = {
    sentence: clause
    for shop in SHOPS.values()
    for _, sentence, clause in shop["sections"]
}


class FakePdfPipeline(Pipeline):
    """A converter that answers with a document instead of reading one.

    Docling's layout and table models are a download and a GPU's worth of
    work; what parsing does with a converted document is tested against
    real Docling documents in the unit layer.
    """

    media_types: ClassVar[tuple[str, ...]] = ("application/pdf",)

    def convert(self, source: SourceDocument) -> Conversion:
        """Builds a two-paragraph document, or a shop's terms where the file says so."""
        from docling_core.types.doc.document import DoclingDocument
        from docling_core.types.doc.labels import DocItemLabel

        document = DoclingDocument(name=source.sha256[:12])
        shop = next((one for key, one in SHOPS.items() if key in source.data), None)
        if shop is None:
            document.add_title(text="The Device")
            for paragraph in PARAGRAPHS:
                document.add_text(label=DocItemLabel.TEXT, text=paragraph)
        else:
            document.add_title(text=shop["title"])
            for heading, sentence, _ in shop["sections"]:
                document.add_heading(text=heading, level=1)
                document.add_text(label=DocItemLabel.TEXT, text=sentence)
        return Conversion(document=document, confidence=0.9, confidence_low=None)


def _statement(line: str) -> str:
    """The claim the stub reads out of one numbered sentence."""
    for sentence, clause in CLAUSES.items():
        if sentence in line:
            return clause
    return (
        "The device weighs 4 kg."
        if "4 kg" in line
        else "The device arrives in March 2026."
    )


class StubModel:
    """A served model that reads one claim out of each sentence."""

    model = "ollama/stub"
    temperature = 0.0

    def __init__(self) -> None:
        """Initialises the call counter."""
        self.calls = 0

    def answer(self, *, system: str, user: str, shape, prompt_version=None):
        """Answers with one fact per numbered sentence in the excerpt."""
        self.calls += 1
        numbered = [line for line in user.splitlines() if line.startswith("[")]
        return shape(
            facts=[
                {
                    "sentences": [index],
                    "statement": _statement(line),
                }
                for index, line in enumerate(numbered)
            ]
        )


#: What question generation is configured with here. Constructed rather than
#: loaded: these are the run's settings, not the deployment's.
QUESTIONS = QuestionSettings(
    per_topic=4,
    sample_size=4,
    samples_per_passage=1,
    fact_kinds=("atomic",),
    type_mix={"factoid": 3, "reason": 1},
    difficulty_mix={"easy": 1, "medium": 1},
    followup_types=("condition",),
    unanswerable_share=0.25,
    followup_share=0.5,
    max_followups=2,
    # Nothing is asked twice: this drives the whole pipeline with stub
    # models, and a retry would double the calls the run is counted by.
    retries=0,
    answer_chars={"value": (1, 80), "list": (3, 300), "explanation": (20, 600)},
    explanation_chars=(150, 900),
    answer_coverage=0.0,
    party_density=0.0,
    meets_floor=0.0,
    answer_overlap=0.6,
    off_topic_overlap=0.3,
    elsewhere_passages=0,
    long_answer_chars=60,
    boilerplate_cosine=0.0,
    duplicate_cosine=0.93,
    duplicate_floor=0.75,
    release_size=0,
    release_unanswerable=0.1,
    release_difficulty={"easy": 1, "medium": 1, "hard": 1},
    embedding_model="stub",
    max_tokens=512,
    model=None,
    verifier_model="ollama/verifier",
)


#: A four-digit year, which is what makes the proper noun before it read as
#: an author. The stub keeps out of the way of `leaks_source` with it.
_YEAR = re.compile(r"\b(1[89]|20)\d{2}\b")

#: What the stub explains its answer with. Every figure and name in it -
#: 4 kg, 12 hours, March 2026, Hamburg - is in PARAGRAPHS above, because the
#: gate refuses an explanation asserting one the passages do not carry. Long
#: enough to clear the floor in QUESTIONS, which is what stops an
#: explanation being the key restated.
EXPLANATION = (
    "The device weighs 4 kg, which is the figure recorded for it against a "
    "run time of 12 hours. The weight is a property of the device itself "
    "rather than of a particular delivery, so it does not change with the "
    "March 2026 arrival from the Hamburg plant. It is the figure to quote "
    "when the device has to be carried or mounted."
)


class StubWriter:
    """A served model that turns each group of facts into one question."""

    model = "ollama/stub"
    temperature = 0.0

    def __init__(self) -> None:
        """Initialises the call counter."""
        self.calls = 0

    def answer(self, *, system: str, user: str, shape, prompt_version=None):
        """Writes a question naming the facts it was given.

        The FACTS block alone, never the PASSAGE one below it. Both are
        rendered as `[n] ...` lines, so taking every line that opens with a
        bracket spliced the passage and its heading trail into the question
        - and `leaks_source` refused the result, correctly. A question
        carrying `Under: The Device` has named where its answer lives.

        The `[n]` labels come off too, for the same gate and a second
        reason: a bracketed number is how a bibliography writes a
        reference, so `[1]` in a question reads as a citation whatever it
        was numbering.

        And a fact carrying a year is not quoted at all. `March 2026` is a
        proper noun with a year directly after it, which is how `Beck 2003`
        is written, so a question repeating it has cited an author. The
        gate is right about the shape and wrong about this fact; a stub
        that has to quote something quotes one of the others.
        """
        self.calls += 1
        facts, _, _ = user.partition("PASSAGE(S)")
        quotable = [
            line.partition("] ")[2]
            for line in facts.splitlines()
            if line.startswith("[") and not _YEAR.search(line)
        ]
        asked = (
            f"Which of these does the device do: {' '.join(quotable)}?"
            if quotable
            else "What does the device do?"
        )
        if "answer" in shape.model_fields:
            return shape(question=asked, answer="4 kg", explanation=EXPLANATION)
        return shape(question=f"On Mars, {asked}")


class StubVerifier:
    """A second model that finds the answer only in an answerable question."""

    model = "ollama/stub-verifier"
    temperature = 0.0

    def __init__(self) -> None:
        """Initialises the call counter."""
        self.calls = 0

    def answer(self, *, system: str, user: str, shape, prompt_version=None):
        """Recovers 4 kg unless the question was perturbed onto Mars."""
        self.calls += 1
        found = "Mars" not in user
        return shape(
            in_passage=found,
            answer="4 kg" if found else "",
            subject="the device",
        )


class StubEmbedder:
    """A deterministic embedding, so the dedup gate still behaves.

    The real one is 2.2 GB of weights, and what the gate needs of it here is
    only that the same question embeds the same way twice and a different one
    does not. A one-hot vector off the digest gives exactly that, and gives
    it without the download.
    """

    def embed(self, text: str) -> list[float]:
        """Answers with the unit vector this text always gets."""
        import hashlib

        vector = [0.0] * 1024
        vector[int.from_bytes(hashlib.sha256(text.encode()).digest()[:2]) % 1024] = 1.0
        return vector


def fit_topics(engine) -> int:
    """Writes the topic a fit would have produced, over every passage.

    Stood in for like the converter and the model, and for a reason of its
    own: the frequency filter needs more documents than this fixture has, so
    a real fit over two passages produces no vocabulary and no topics. What
    the stage after it needs is a topic and the memberships, which is what
    this writes.
    """
    with engine.begin() as connection:
        topic_id = connection.execute(
            text(
                "INSERT INTO topics (language, topic_index, top_terms, status, "
                "requested_at, fitted_at) VALUES ('en', 0, ARRAY['device'], "
                "'modelled', now(), now()) RETURNING id"
            )
        ).scalar_one()
        connection.execute(
            text(
                "INSERT INTO passage_topics (passage_id, topic_id, weight) "
                "SELECT id, :topic, 0.9 FROM passages"
            ),
            {"topic": topic_id},
        )
    return topic_id


@pytest.fixture(scope="module")
def chunker():
    """The real passage builder, or a skip when its tokenizer is not here.

    It is a HuggingFace tokenizer, downloaded on first use. A machine with
    no cache and no network cannot build one, and that is a reason to skip
    this rather than to stand it in for.
    """
    from preprocessing.chunking.passages import PassageBuilder

    try:
        return PassageBuilder(
            embedding_model="intfloat/multilingual-e5-large",
            max_tokens=512,
            merge_peers=True,
        )
    except Exception as exc:  # noqa: BLE001 - any failure is the same answer
        weights.unavailable(f"the chunker's tokenizer is unavailable: {exc}")


@pytest.fixture
def pipeline(database, buckets, chunker):
    """Every stage, wired as its factory wires it."""
    return _pipeline_with(chunker, StubModel())


def _pipeline_with(chunker, model) -> dict:
    """Builds the stages around one served model."""
    from blob_store.s3 import DocumentsBucket, ParsedBucket

    return {
        "ingest": IngestService(
            repository=DocumentRepository(),
            store=DocumentsBucket(),
            max_file_size_bytes=1024 * 1024,
            allowed_media_types=("application/pdf",),
            pipeline_version="0.0.1",
        ),
        "parse": ParsingService(
            repository=ParseQueue(),
            documents=DocumentsBucket(),
            parsed=ParsedBucket(),
            pipelines=PipelineRegistry((FakePdfPipeline(),)),
            analyser=DocumentAnalyser(),
            ocr_char_threshold=100,
            min_confidence=0.5,
        ),
        "chunk": ChunkingService(
            repository=ChunkQueue(), parsed=ParsedBucket(), builder=chunker
        ),
        "extract": ExtractionService(
            repository=PassageQueue(),
            extractors=ExtractorRegistry(
                extractors=(TableExtractor(),), default=LlmExtractor(model)
            ),
            checker=FactChecker(),
        ),
        "questions": QuestionGenerationService(
            repository=QuestionQueue(),
            writer=QuestionWriter(StubWriter()),
            checker=QuestionChecker(
                embedder=StubEmbedder(),
                verifier=Verifier(StubVerifier()),
                nearest=QuestionCatalog().nearest,
                threshold=QUESTIONS.duplicate_cosine,
                floor=QUESTIONS.duplicate_floor,
            ),
            settings=QUESTIONS,
        ),
        "model": model,
    }


def counts(engine) -> dict[str, int]:
    """How many rows each table holds."""
    with engine.connect() as connection:
        return {
            table: connection.execute(
                text(f"SELECT count(*) FROM {table}")
            ).scalar_one()
            for table in ("documents", "passages", "facts", "questions")
        }


def statuses(engine) -> tuple[str, str]:
    """The parse and chunk status of the one document."""
    with engine.connect() as connection:
        return connection.execute(
            text("SELECT parse_status, chunk_status FROM documents")
        ).one()


def test_nothing_runs_until_something_asks_it_to(pipeline, engine) -> None:
    """A row arrives `new` and no worker looks at it."""
    pipeline["ingest"].ingest(UploadedFile("report.pdf", PDF))

    assert pipeline["parse"].drain() == 0
    assert statuses(engine) == (Status.NEW, Status.NEW)


def test_a_document_becomes_facts(pipeline, engine) -> None:
    """Upload, parse, chunk, extract, each stage claiming from the last."""
    stored = pipeline["ingest"].ingest(UploadedFile("report.pdf", PDF))
    assert stored.sha256

    ParseQueue().start()
    assert pipeline["parse"].drain() == 1
    assert statuses(engine)[0] == Status.PARSED

    ChunkQueue().start()
    assert pipeline["chunk"].drain() == 1
    assert statuses(engine)[1] == Status.CHUNKED

    PassageQueue().start()
    read = pipeline["extract"].drain()

    assert read > 0
    held = counts(engine)
    assert held["documents"] == 1
    assert held["passages"] > 0
    assert held["facts"] > 0


def test_every_stored_fact_cites_a_sentence_of_its_passage(pipeline, engine) -> None:
    """The citation is an index, so it has to index something.

    One assertion per link row rather than one per fact: a claim resting on
    two passages has to resolve in both, and the joined spans have to be the
    evidence the search index holds.
    """
    _run(pipeline)

    with engine.connect() as connection:
        rows = connection.execute(
            text(
                "SELECT fp.fact_id, f.evidence_text, "
                "substring(p.text FROM fp.evidence_start + 1 "
                "          FOR fp.evidence_end - fp.evidence_start) AS cited "
                "FROM fact_passages fp "
                "JOIN facts f ON f.id = fp.fact_id "
                "JOIN passages p ON p.id = fp.passage_id "
                "WHERE fp.sentence_ids IS NOT NULL "
                "ORDER BY fp.fact_id, fp.position"
            )
        ).all()

    assert rows
    joined: dict[int, list[str]] = {}
    evidence: dict[int, str] = {}
    for fact_id, stored, cited in rows:
        joined.setdefault(fact_id, []).append(cited)
        evidence[fact_id] = stored
    for fact_id, spans in joined.items():
        assert "\n".join(spans) == evidence[fact_id]


def test_the_model_is_asked_once_per_passage_that_carries_a_claim(
    pipeline, engine
) -> None:
    """Once each, over a corpus whose every passage carries one.

    The skip a heading or a caption takes is test_passage_gate.py's: this
    corpus chunks into passages that all carry a claim.
    """
    _run(pipeline)

    with engine.connect() as connection:
        passages = connection.execute(
            text("SELECT count(*) FROM passages")
        ).scalar_one()

    assert pipeline["model"].calls == passages > 0


def test_a_second_run_of_a_finished_stage_does_nothing(pipeline) -> None:
    """Draining an empty queue is not an error."""
    _run(pipeline)

    assert pipeline["parse"].drain() == 0
    assert pipeline["chunk"].drain() == 0
    assert pipeline["extract"].drain() == 0


def test_dropping_the_derived_data_lets_the_run_repeat(pipeline, engine) -> None:
    """Chunking returns to the queue without re-parsing."""
    from blob_store.s3 import ArchiveBucket, DocumentsBucket, ParsedBucket
    from ingestion.removal import RemovalService

    sha = _run(pipeline)
    before = counts(engine)["passages"]

    RemovalService(
        repository=DocumentRepository(),
        store=DocumentsBucket(),
        parsed=ParsedBucket(),
        archive=ArchiveBucket(),
    ).delete_derived(sha)
    assert counts(engine)["passages"] == 0

    ChunkQueue().start()
    assert pipeline["chunk"].drain() == 1

    assert counts(engine)["passages"] == before
    assert statuses(engine)[0] == Status.PARSED, "it did not need re-parsing"


def test_a_model_that_will_not_answer_fails_the_passage(
    database, buckets, chunker, engine
) -> None:
    """The row records the reason and can be returned to the queue."""

    class Refusing:
        """A served model that is never available."""

        model, temperature = "ollama/stub", 0.0

        def answer(self, **_: object):
            """Refuses every ask."""
            from llm.client import ModelUnavailable

            raise ModelUnavailable("Timeout: the model did not answer")

    built = _pipeline_with(chunker, Refusing())
    _run(built)

    with engine.connect() as connection:
        failed = connection.execute(
            text(
                "SELECT extract_status, extract_error FROM passages "
                "WHERE extract_status = 'failed'"
            )
        ).all()

    assert failed, "no passage was marked failed"
    assert all("did not answer" in error for _, error in failed)
    assert PassageQueue().retry() == len(failed), "a failure must be retryable"


def _run(pipeline) -> str:
    """Runs one document through every stage, and answers with its digest."""
    stored = pipeline["ingest"].ingest(UploadedFile("report.pdf", PDF))
    ParseQueue().start()
    pipeline["parse"].drain()
    ChunkQueue().start()
    pipeline["chunk"].drain()
    PassageQueue().start()
    pipeline["extract"].drain()
    assert stored.sha256
    return stored.sha256


def test_the_digest_survives_every_stage(pipeline, engine) -> None:
    """One identity from the upload to the fact."""
    sha = _run(pipeline)

    with engine.connect() as connection:
        held = (
            connection.execute(
                text(
                    "SELECT DISTINCT p.doc_sha256 FROM passages p "
                    "JOIN fact_passages fp ON fp.passage_id = p.id"
                )
            )
            .scalars()
            .all()
        )

    assert held == [sha]


def _through_questions(pipeline, engine) -> int:
    """Runs one document all the way to its questions, and answers with the topic."""
    _run(pipeline)
    topic_id = fit_topics(engine)
    QuestionQueue().start()
    pipeline["questions"].drain()
    return topic_id


def test_a_document_becomes_questions(pipeline, engine) -> None:
    """The sixth stage, claiming from a topic the way the rest claim from a row.

    One of them has to be ACCEPTED, not merely written. A rejected question
    is stored like any other, so a run whose every question a gate refused
    counts the same here as a run that worked - and one did: the stub wrote
    a question carrying its passage's heading trail, `leaks_source` refused
    it correctly, and this test went on passing on the row it left behind.
    """
    topic_id = _through_questions(pipeline, engine)

    held = counts(engine)
    assert held["questions"] > 0
    with engine.connect() as connection:
        assert connection.execute(
            text("SELECT count(*) FROM questions WHERE status = 'accepted'")
        ).scalar(), "every question written was refused by a gate"
        assert (
            connection.execute(
                text("SELECT question_status FROM topics WHERE id = :id"),
                {"id": topic_id},
            ).scalar()
            == Status.GENERATED
        )


def test_an_accepted_question_carries_both_of_its_answers(pipeline, engine) -> None:
    """The key that is matched, and the reading beside it.

    Two columns because two gates: recoverability compares the target and
    divides by its own lemmas, so the answer that teaches cannot also be
    the answer that is matched.
    """
    _through_questions(pipeline, engine)

    with engine.connect() as connection:
        answers = connection.execute(
            text(
                "SELECT target_answer, answer_explanation FROM questions "
                "WHERE status = 'accepted' AND answerable"
            )
        ).all()

    assert answers, "no answerable question was accepted"
    for target, explanation in answers:
        assert target, "an accepted answerable question carries no key"
        assert explanation, "an accepted answerable question carries no reading"
        assert len(explanation) > len(target)


def test_nothing_generates_until_something_asks_it_to(pipeline, engine) -> None:
    """A fitted topic arrives `new`, like every other row in the pipeline."""
    _run(pipeline)
    fit_topics(engine)

    assert pipeline["questions"].drain() == 0
    assert counts(engine)["questions"] == 0


def test_every_question_can_be_traced_back_to_a_passage(pipeline, engine) -> None:
    """Through its facts, which is the only route there is."""
    _through_questions(pipeline, engine)

    with engine.connect() as connection:
        orphans = connection.execute(
            text(
                "SELECT count(*) FROM questions q WHERE NOT EXISTS ("
                "  SELECT 1 FROM question_facts qf "
                "  JOIN fact_passages fp ON fp.fact_id = qf.fact_id "
                "  JOIN passages p ON p.id = fp.passage_id "
                "  WHERE qf.question_id = q.id)"
            )
        ).scalar_one()

    assert orphans == 0


def _resting(engine) -> list:
    """Each question, its status and reason, and the facts it rests on."""
    with engine.connect() as connection:
        return list(
            connection.execute(
                text(
                    "SELECT q.id, q.status, q.rejected_reason, "
                    "array_remove(array_agg(qf.fact_id ORDER BY qf.fact_id), NULL) "
                    "FROM questions q "
                    "LEFT JOIN question_facts qf ON qf.question_id = q.id "
                    "GROUP BY q.id ORDER BY q.id"
                )
            ).all()
        )


class _Rewording(StubModel):
    """The stub, reading each claim in other words than it did the first time."""

    def answer(self, *, system: str, user: str, shape, prompt_version=None):
        """One fact per numbered sentence, none of them as first read."""
        self.calls += 1
        numbered = [line for line in user.splitlines() if line.startswith("[")]
        return shape(
            facts=[
                {"sentences": [index], "statement": f"Recorded: {_statement(line)}"}
                for index, line in enumerate(numbered)
            ]
        )


def test_re_extracting_the_corpus_keeps_its_questions(
    pipeline, chunker, engine
) -> None:
    """A re-extraction never deletes a question; it rejects one whose fact went.

    Read again as it was, a fact keeps its id and the question resting on it
    stays as it was. Read otherwise, the fact goes and its question is
    rejected as `source_changed` and kept, with its id, where it once went
    with the fact.
    """
    _through_questions(pipeline, engine)
    before = _resting(engine)
    assert before

    PassageQueue().reset()
    pipeline["extract"].drain()

    assert _resting(engine) == before, "a fact read again as it was moved a question"

    PassageQueue().reset()
    _pipeline_with(chunker, _Rewording())["extract"].drain()

    after = _resting(engine)
    assert [one[0] for one in after] == [one[0] for one in before], "a question went"
    assert all(one[1:3] == ("rejected", "source_changed") for one in after), after


def test_re_judging_a_fact_leaves_the_question_for_the_re_check_to_find(
    pipeline, engine
) -> None:
    """The quiet path: nothing is deleted, and the question is now wrong.

    `extract-revalidate` updates a fact in place, so a question resting on
    one that no longer passes its checks survives with nothing to say so.
    That is what questions-reverify is for.
    """
    _through_questions(pipeline, engine)
    before = counts(engine)["questions"]

    with engine.begin() as connection:
        connection.execute(
            text(
                "UPDATE facts SET validated = false, rejection_code = 'copied' "
                "WHERE id = (SELECT min(fact_id) FROM question_facts)"
            )
        )

    assert counts(engine)["questions"] == before, "nothing should have been deleted"
    assert reverify(QuestionCatalog(), QUESTIONS) > 0

    with engine.connect() as connection:
        reasons = connection.scalars(
            text("SELECT rejected_reason FROM questions WHERE status = 'rejected'")
        ).all()
    assert "source_changed" in reasons


def test_a_re_check_never_overturns_a_person(pipeline, engine) -> None:
    """Accepting is a person's, and a nightly re-check must not undo it.

    One fact is re-judged too, so the re-check has a reason it could write.
    """
    _through_questions(pipeline, engine)
    with engine.begin() as connection:
        connection.execute(
            text("UPDATE questions SET status = 'rejected', rejected_reason = NULL")
        )
        connection.execute(
            text(
                "UPDATE facts SET validated = false, rejection_code = 'copied' "
                "WHERE id = (SELECT min(fact_id) FROM question_facts)"
            )
        )

    reverify(QuestionCatalog(), QUESTIONS)

    with engine.connect() as connection:
        states = set(
            connection.execute(
                text("SELECT DISTINCT status, rejected_reason FROM questions")
            ).all()
        )
    assert states == {("rejected", None)}, "a re-check rewrote a person's verdict"


# ── Through the bridges, the release, the judge and the workbook ───────────


class StubBridge:
    """A served model that joins the first fact of each of two excerpts."""

    model = "ollama/stub-bridge"
    temperature = 0.0

    def __init__(self) -> None:
        """Initialises the call counter."""
        self.calls = 0

    def answer(self, *, system: str, user: str, shape, prompt_version=None):
        """Joins the facts the excerpts list, in the language they are written in."""
        self.calls += 1
        listed = [
            line.partition("] ")[2] for line in user.splitlines() if "[F0]" in line
        ]
        german = "German" in user
        first, second = (one.rstrip(".") for one in listed[:2])
        joined = (
            f"{first}, und {second[0].lower()}{second[1:]}."
            if german
            else f"{first}, and {second[0].lower()}{second[1:]}."
        )
        return shape(
            relation="consequence",
            facts=[{"excerpt": 0, "fact": 0}, {"excerpt": 1, "fact": 0}],
            statement=joined,
            question="Wie lange dauert beides?" if german else "How long is each step?",
            answer="2 Werktage" if german else "2 working days",
        )


class StubJudge:
    """An independent judge that approves everything it is asked about."""

    model = "ollama/stub-judge"

    def judge(self, artifact):
        """One approving judgement for each metric the kind is judged on."""
        from assessment.models import Judgement
        from assessment.templates import _CHOICES, metrics_of

        found = []
        for metric in metrics_of(artifact.kind):
            scores, good = _CHOICES[metric]
            found.append(
                Judgement(
                    metric=metric,
                    label=good,
                    score=scores[good],
                    approved=True,
                    explanation="approved by the stand-in",
                )
            )
        return tuple(found)

    def due(self, artifact) -> int:
        """Every metric the kind declares."""
        from assessment.templates import metrics_of

        return len(metrics_of(artifact.kind))


def fit_topics_by_language(engine) -> dict[str, int]:
    """One topic a language, holding every passage of that language.

    The stand-in for the fit, as `fit_topics` is, but per language: a
    German passage joins a German topic, which is what the bridge pass and
    the writer both read.
    """
    topics: dict[str, int] = {}
    with engine.begin() as connection:
        for (language,) in connection.execute(
            text("SELECT DISTINCT language FROM passages ORDER BY 1")
        ).all():
            topic = connection.execute(
                text(
                    "INSERT INTO topics (language, topic_index, top_terms, status, "
                    "requested_at, fitted_at) VALUES (:language, 0, ARRAY['shop'], "
                    "'modelled', now(), now()) RETURNING id"
                ),
                {"language": language},
            ).scalar_one()
            connection.execute(
                text(
                    "INSERT INTO passage_topics (passage_id, topic_id, weight) "
                    "SELECT id, :topic, 0.9 FROM passages WHERE language = :language"
                ),
                {"topic": topic, "language": language},
            )
            topics[language] = topic
    return topics


def _everything(pipeline, engine) -> dict:
    """The device document and two shops' terms, through to a drawn release.

    Upload, parse, chunk and extract; a topic per language; the bridge pass;
    the questions; the release draw; and the judge. What each stage leaves is
    what the next one claims, which is the handover this file is about.
    """
    from dataclasses import replace

    from assessment.repository import AssessmentQueue
    from assessment.service import AssessmentService
    from extraction.extractors.bridge import BridgeExtractor
    from extraction.repository import FactCatalog
    from extraction.service import bridge
    from question_generation.service import balance

    for name, marker in (
        ("device.pdf", b""),
        ("shop.pdf", b"%shop-en"),
        ("laden.pdf", b"%shop-de"),
    ):
        pipeline["ingest"].ingest(UploadedFile(name, PDF + marker))
    ParseQueue().start()
    pipeline["parse"].drain()
    ChunkQueue().start()
    pipeline["chunk"].drain()
    PassageQueue().start()
    pipeline["extract"].drain()

    topics = fit_topics_by_language(engine)
    model = StubBridge()
    bridged = bridge(
        FactCatalog(),
        BridgeExtractor(model),
        FactChecker(),
        groups_per_topic=2,
        size=2,
    )

    QuestionQueue().start()
    pipeline["questions"].drain()
    settings = replace(QUESTIONS, release_size=1)
    released = balance(QuestionCatalog(), settings, None, Verifier(StubVerifier()))

    queue = AssessmentQueue(kinds=("question",), sample=0, stratify=False)
    enrolled = queue.enrol()
    queue.start()
    judged = AssessmentService(queue, StubJudge()).drain()
    return {
        "topics": topics,
        "bridged": bridged,
        "bridge_calls": model.calls,
        "released": released,
        "enrolled": enrolled,
        "judged": judged,
    }


def test_the_run_goes_on_through_bridges_the_release_and_the_judge(
    pipeline, engine
) -> None:
    """Each stage after questions finds what the one before it left.

    A German document and an English one, both a shop's terms: a bridge is
    written in each language, and what is drawn into the release is judged.
    """
    ran = _everything(pipeline, engine)

    assert set(ran["topics"]) >= {"de", "en"}
    assert ran["bridge_calls"] >= 2, "a link was offered in each language"
    assert ran["bridged"] >= 2, "and each was stored"
    with engine.connect() as connection:
        bridges = (
            connection.execute(
                text(
                    "SELECT DISTINCT p.language FROM facts f "
                    "JOIN fact_passages fp ON fp.fact_id = f.id "
                    "JOIN passages p ON p.id = fp.passage_id WHERE f.kind = 'bridge'"
                )
            )
            .scalars()
            .all()
        )
    assert sorted(bridges) == ["de", "en"]
    assert ran["released"] > 0, "nothing was drawn into the release"
    assert ran["enrolled"] > 0 and ran["judged"] == ran["enrolled"]


def test_the_workbook_holds_the_release_its_judge_and_its_summary(
    pipeline, engine, tmp_path
) -> None:
    """The last thing a person takes away: the released questions, read back."""
    from openpyxl import load_workbook

    from question_generation.export import main

    _everything(pipeline, engine)
    out = tmp_path / "release.xlsx"

    assert main(["--out", str(out), "--released"]) == 0

    book = load_workbook(out, read_only=True)
    assert book.sheetnames == ["Questions", "Citations", "Assessment", "Summary"]
    head, *released = list(book["Questions"].iter_rows(values_only=True))
    column = {name: at for at, name in enumerate(head)}
    assert released, "the sheet holds a header and no released question"
    for row in released:
        assert row[column["Released"]] is True
        assert row[column["Status"]] == "accepted"
        assert row[column["Language"]] in {"de", "en"}
        assert row[column["Judge"]] == "approved"
    assert len(list(book["Assessment"].iter_rows(values_only=True))) > 1
    summary = dict(list(book["Summary"].iter_rows(values_only=True))[1:])
    assert summary["Questions in this workbook"] == len(released)
    assert summary["Judged by a model"] == len(released)
    assert summary["Judge approved"] == len(released)
