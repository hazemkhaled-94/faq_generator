"""The workbook as one moment, while the question worker goes on writing.

`GET /questions/export` reads the rows, then the citations, then the quality
report, as three statements in three sessions. A run writes questions all
the while, so a question written between the first read and the others is
a moment a real export can land in. This writes one there, as the worker
would, and reads what the workbook then says about itself.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from seed import digest, document, fact, link, passage, question
from sqlalchemy.orm import Session

pytestmark = pytest.mark.integration


@pytest.fixture
def exported(client, engine, monkeypatch):
    """A one-question corpus, exported while a second question is written."""
    from api.routes import questions as routes

    with Session(engine) as session:
        session.add(document(digest("shop"), title="Rückgabe"))
        session.flush()
        held = passage(digest("shop"), text="Eine Rücksendung ist 30 Tage frei.")
        session.add(held)
        session.flush()
        cited = fact(held.id, statement="Eine Rücksendung ist 30 Tage frei.")
        session.add(cited)
        session.flush()
        first = question(question_text="Wie lange ist eine Rücksendung frei?")
        session.add(first)
        session.flush()
        session.add(link(first.id, cited.id))
        session.commit()
        fact_id = cited.id

    catalog = routes.question_catalog
    listed = catalog.page

    def page_then_a_write(*args, **kwargs):
        """The listing, and then the worker writing its next question."""
        found = listed(*args, **kwargs)
        with Session(engine) as session:
            later = question(question_text="Was kostet eine Rücksendung danach?")
            session.add(later)
            session.flush()
            session.add(link(later.id, fact_id))
            session.commit()
        return found

    monkeypatch.setattr(catalog, "page", page_then_a_write)
    answer = client.get("/questions/export")
    assert answer.status_code == 200
    return load_workbook(io.BytesIO(answer.content))


def test_the_workbook_describes_the_rows_it_holds(exported) -> None:
    """The Summary counts the Questions sheet, and every citation has its row."""
    head, *rows = exported["Questions"].iter_rows(values_only=True)
    held = {row[head.index("ID")] for row in rows}
    cites = {
        row[0] for row in list(exported["Citations"].iter_rows(values_only=True))[1:]
    }
    summary = {
        row[0]: row[1]
        for row in list(exported["Summary"].iter_rows(values_only=True))[1:]
    }

    assert cites <= held
    assert summary["Matching the filter"] == summary["Questions in this workbook"]
