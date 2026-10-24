"""Every page, drawn from what the real API answers rather than a script.

`tests/frontend` drives each view against scripted answers, which say what a
page does with an answer somebody wrote down. These say what it does with
the answers the API actually gives: an empty database, a small corpus, and
the workbook the Questions page's Build button asks for.
"""

from __future__ import annotations

import io
import os
from pathlib import Path

import pytest
import requests
from openpyxl import load_workbook
from seed import digest, document, fact, fitted, link, membership, passage, question
from sqlalchemy.orm import Session

pytestmark = [pytest.mark.integration, pytest.mark.frontend]

VIEWS = Path(__file__).resolve().parents[3] / "frontend/views"

#: Every page the sidebar offers.
PAGES = (
    "upload",
    "documents",
    "passages",
    "facts",
    "topics",
    "questions",
    "assessment",
    "health",
)

#: The client factories a view reaches for, by the class each one builds.
_CLIENTS = {
    "catalog_api": "CatalogApi",
    "settings_api": "SettingsApi",
    "upload_api": "UploadApi",
    "health_api": "HealthApi",
    "pipeline_api": "PipelineApi",
    "lineage_api": "LineageApi",
}

os.environ.setdefault("BACKEND_URL", "http://testserver")
os.environ.setdefault("PAGE_SIZE", "25")


class _Wire:
    """A `requests.Session` that answers from the app in this process."""

    def __init__(self, client) -> None:
        """Wraps the TestClient every route is reached through."""
        self._client = client

    def request(self, method: str, url: str, timeout=None, **kwargs):
        """Sends one request and hands back what `requests` would have."""
        del timeout
        got = self._client.request(method, url, **kwargs)
        answer = requests.Response()
        answer.status_code = got.status_code
        answer._content = got.content
        answer.headers = requests.structures.CaseInsensitiveDict(got.headers)
        answer.url = url
        answer.reason = got.reason_phrase
        answer.encoding = got.encoding
        return answer

    def get(self, url: str, **kwargs):
        """The one verb a client sends past `_request`."""
        return self.request("GET", url, **kwargs)


@pytest.fixture
def drawn(client, monkeypatch):
    """Runs one view against the real API, and returns the finished app."""
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    import lib.backend as module

    wire = _Wire(client)
    for factory, built in _CLIENTS.items():
        made = getattr(module, built)
        monkeypatch.setattr(
            module, factory, lambda made=made: made("http://testserver", wire)
        )

    def run(name: str):
        st.cache_data.clear()
        return AppTest.from_file(str(VIEWS / f"{name}.py"), default_timeout=60).run()

    yield run
    st.cache_data.clear()


def _failures(app) -> list[str]:
    """What the page raised or drew as a page-wide failure."""
    return [str(one.value) for one in app.exception] + [
        str(one.value) for one in app.error if "Failed to load" in str(one.value)
    ]


@pytest.mark.parametrize("name", PAGES)
def test_every_page_renders_on_an_empty_database(drawn, name) -> None:
    """A fresh deployment: no document, no topic, no question, nothing judged."""
    app = drawn(name)

    assert _failures(app) == []


@pytest.fixture
def corpus(client, engine):
    """One bank document in German, its passages, facts, a topic and questions."""
    with Session(engine) as session:
        session.add(
            document(
                digest("bank"),
                language="de",
                title="Preis- und Leistungsverzeichnis",
                parse_status="parsed",
                chunk_status="chunked",
            )
        )
        subject = fitted(0, language="de", top_terms=["konto", "gebühr"], label="Konto")
        session.add(subject)
        session.flush()
        first = passage(
            digest("bank"),
            ordinal=1,
            text="Die Kontoführung kostet 4,90 EUR im Monat.",
            language="de",
        )
        second = passage(
            digest("bank"),
            ordinal=2,
            text="Eine Ersatzkarte kostet 15 EUR.",
            language="de",
        )
        session.add_all([first, second])
        session.flush()
        session.add_all(membership(one.id, subject.id) for one in (first, second))
        facts = [
            fact(first.id, statement="Die Kontoführung kostet 4,90 EUR im Monat."),
            fact(second.id, statement="Eine Ersatzkarte kostet 15 EUR."),
        ]
        session.add_all(facts)
        session.flush()
        asked = [
            question(
                question_text="Was kostet die Kontoführung im Monat?",
                target_answer="4,90 EUR",
                language="de",
                status="accepted",
                difficulty="easy",
                passage_scope="single_passage",
                document_scope="single_document",
                topic_scope="single_topic",
                question_type="factoid",
                answer_form="value",
                cognitive_level="recall",
                kind="single",
                answer_chars=8,
            ),
            question(
                question_text="Was kostet die Kontoführung, und was eine Ersatzkarte?",
                target_answer="4,90 EUR; 15 EUR",
                language="de",
                status="draft",
                difficulty="medium",
                passage_scope="multi_passage",
                document_scope="single_document",
                topic_scope="single_topic",
                kind="compound",
                answer_chars=16,
            ),
        ]
        session.add_all(asked)
        session.flush()
        session.add_all(
            [
                link(asked[0].id, facts[0].id),
                link(asked[1].id, facts[0].id),
                link(asked[1].id, facts[1].id),
            ]
        )
        session.commit()
        return client, {"topic": subject.id, "questions": [one.id for one in asked]}


@pytest.mark.parametrize("name", PAGES)
def test_every_page_renders_on_a_small_corpus(drawn, corpus, name) -> None:
    """The same pages over a document, its passages, facts, a topic and questions."""
    app = drawn(name)

    assert _failures(app) == []


def test_the_questions_page_lists_what_the_api_holds(drawn, corpus) -> None:
    """Both questions, under the total the API counted."""
    app = drawn("questions")
    shown = " ".join(one.value.to_string() for one in app.dataframe)

    assert "Was kostet die Kontoführung im Monat?" in shown
    assert "compound" in shown


def test_the_workbook_holds_the_kind_of_every_question(corpus) -> None:
    """The export the Build button asks for, read back as Excel reads it."""
    client, _ = corpus

    answer = client.get("/questions/export")
    book = load_workbook(io.BytesIO(answer.content))
    rows = list(book["Questions"].iter_rows(values_only=True))
    summary = {
        row[0]: row[1] for row in list(book["Summary"].iter_rows(values_only=True))[1:]
    }

    assert answer.status_code == 200
    assert sorted(row[rows[0].index("Kind")] for row in rows[1:]) == [
        "compound",
        "single",
    ]
    assert summary["kind: compound"] == 1
    assert summary["Questions in this workbook"] == 2


def test_a_topic_named_with_a_control_character_still_exports(corpus) -> None:
    """The name is either refused at the edge or written without the character."""
    client, held = corpus

    named = client.patch(f"/topics/{held['topic']}", json={"label": "Konto\x0bgebühr"})
    exported = client.get("/questions/export")

    assert named.status_code in (200, 422)
    assert exported.status_code == 200
