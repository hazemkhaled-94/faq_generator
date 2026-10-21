"""How the workbook names a document, on the two sheets that name one.

A PDF with no title item is common: parsing stores None, and every reader
falls back to the digest. The Questions sheet's Documents column and the
Citations sheet's Document column are read side by side by somebody
checking an answer, so they have to name one document the same way.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from seed import digest, document, fact, link, passage, question
from sqlalchemy.orm import Session

pytestmark = pytest.mark.integration


def _sheet(book, name: str) -> list[dict]:
    """One sheet's rows, by heading."""
    head, *rows = book[name].iter_rows(values_only=True)
    return [dict(zip(head, row, strict=True)) for row in rows]


@pytest.fixture
def exported(client, engine):
    """One question over a titled document and an untitled one, exported."""
    with Session(engine) as session:
        session.add_all(
            [
                document(digest("titled"), title="Versandbedingungen"),
                document(digest("untitled"), title=None),
            ]
        )
        session.flush()
        shipped = passage(digest("titled"), text="Ein Paket kommt in 3 Tagen.")
        returned = passage(digest("untitled"), text="Returns are free for 30 days.")
        session.add_all([shipped, returned])
        session.flush()
        facts = [
            fact(shipped.id, statement="Ein Paket kommt in 3 Tagen."),
            fact(returned.id, statement="Returns are free for 30 days."),
        ]
        session.add_all(facts)
        session.flush()
        asked = question(
            question_text="How long does a parcel take, and are returns free?",
            target_answer="3 days; free for 30 days",
            status="accepted",
            document_scope="cross_document",
        )
        session.add(asked)
        session.flush()
        session.add_all(link(asked.id, one.id) for one in facts)
        session.commit()

    answer = client.get("/questions/export")
    assert answer.status_code == 200
    return load_workbook(io.BytesIO(answer.content))


def test_a_titled_document_is_named_by_its_title_on_both_sheets(exported) -> None:
    """The title, never the digest, where the document carries one."""
    (row,) = _sheet(exported, "Questions")
    cited = {one["Document"] for one in _sheet(exported, "Citations")}

    assert "Versandbedingungen" in row["Documents"].split(", ")
    assert "Versandbedingungen" in cited


def test_an_untitled_document_is_named_the_same_way_on_both_sheets(exported) -> None:
    """Every name the Citations sheet gives is one the Questions row gives."""
    (row,) = _sheet(exported, "Questions")
    cited = {one["Document"] for one in _sheet(exported, "Citations")}

    assert cited == set(row["Documents"].split(", "))
